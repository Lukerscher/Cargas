import os
import re
import tempfile
import unittest
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook


os.environ.setdefault("DATABASE_PATH", str(Path(tempfile.gettempdir()) / f"almoxcontrol-tests-{os.getpid()}.sqlite"))

import app as application
from app import Category, ConferenceItem, Invoice, Load, Movement, Product, Supplier, User, db


class AlmoxControlTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        application.app.config.update(TESTING=True, WTF_CSRF_ENABLED=True)

    def setUp(self):
        with application.app.app_context():
            db.drop_all()
            db.create_all()
            self.admin = User(name="Ana Administradora", email="ana@example.com", role="administrador")
            self.admin.set_password("SenhaSegura123")
            self.checker = User(name="Caio Conferente", email="caio@example.com", role="conferente")
            self.checker.set_password("SenhaSegura123")
            category = Category(name="Eletrônicos")
            supplier = Supplier(name="Fornecedor Alfa", document="12345678000100")
            db.session.add_all([self.admin, self.checker, category, supplier])
            db.session.flush()
            self.product = Product(code="PRD-001", barcode="7891234567890", name="Leitor industrial", unit="un", min_stock=2, category=category, supplier=supplier)
            db.session.add(self.product)
            db.session.flush()
            db.session.add(Movement(product=self.product, kind="entrada", quantity=15, reference="NF 100", responsible=self.admin))
            db.session.commit()
            self.product_id = self.product.id
            self.category_id = category.id
            self.supplier_id = supplier.id
            self.admin_id = self.admin.id
        self.client = application.app.test_client()
        page = self.client.get("/login")
        match = re.search(rb'name="csrf_token" value="([^"]+)"', page.data)
        self.assertIsNotNone(match)
        self.csrf = match.group(1).decode()
        response = self.client.post("/login", data={"csrf_token": self.csrf, "email": "ana@example.com", "password": "SenhaSegura123"})
        self.assertEqual(response.status_code, 302)

    def tearDown(self):
        with application.app.app_context():
            db.session.remove()

    def post_json(self, url, body):
        return self.client.post(url, json=body, headers={"X-CSRFToken": self.csrf})

    def test_products_stock_and_shortage_guard(self):
        response = self.post_json("/api/movements", {"product_id": self.product_id, "kind": "saida", "quantity": 16})
        self.assertEqual(response.status_code, 400)
        response = self.post_json("/api/movements", {"product_id": self.product_id, "kind": "saida", "quantity": 3})
        self.assertEqual(response.status_code, 201)
        products = self.client.get("/api/products").get_json()
        self.assertEqual(products[0]["stock"], 12)
        duplicate = self.post_json("/api/products", {"code": "PRD-002", "barcode": "7891234567890", "name": "Duplicado", "unit": "un"})
        self.assertEqual(duplicate.status_code, 400)

    def test_conference_scans_a_product_once_with_total_quantity(self):
        created = self.post_json("/api/loads", {"reference": "LOAD-001", "supplier": "Fornecedor Alfa"})
        load_id = created.get_json()["id"]
        first = self.post_json(f"/api/loads/{load_id}/scan", {"barcode": "7891234567890", "expected_quantity": 20, "quantity": 18})
        self.assertEqual(first.status_code, 201)
        duplicate = self.post_json(f"/api/loads/{load_id}/scan", {"barcode": "7891234567890", "quantity": 1})
        self.assertEqual(duplicate.status_code, 409)
        detail = self.client.get(f"/api/loads/{load_id}").get_json()
        self.assertEqual(len(detail["items"]), 1)
        self.assertEqual(detail["items"][0]["checked"], 18)
        self.assertEqual(detail["items"][0]["expected"], 20)

    def test_load_created_from_invoice_prefills_items_and_scans_against_expected(self):
        xml = b'''<NFe><infNFe Id="NFe92345678901234567890123456789012345678901234"><ide><nNF>228</nNF></ide><emit><xNome>Fornecedor Alfa</xNome></emit><det><prod><cProd>PRD-001</cProd><cEAN>7891234567890</cEAN><xProd>Leitor industrial</xProd><qCom>4</qCom><vUnCom>10</vUnCom></prod></det></infNFe></NFe>'''
        second_xml = b'''<NFe><infNFe Id="NFe10345678901234567890123456789012345678901234"><ide><nNF>229</nNF></ide><emit><xNome>Fornecedor Alfa</xNome></emit><det><prod><cProd>PRD-001</cProd><cEAN>7891234567890</cEAN><xProd>Leitor industrial</xProd><qCom>2</qCom><vUnCom>10</vUnCom></prod></det></infNFe></NFe>'''
        imported = self.client.post("/api/invoices/import", data={"file": (BytesIO(xml), "nf-228.xml")}, headers={"X-CSRFToken": self.csrf}, content_type="multipart/form-data")
        second_imported = self.client.post("/api/invoices/import", data={"file": (BytesIO(second_xml), "nf-229.xml")}, headers={"X-CSRFToken": self.csrf}, content_type="multipart/form-data")
        invoice_id = imported.get_json()["id"]
        second_invoice_id = second_imported.get_json()["id"]
        created = self.post_json("/api/loads", {"reference": "LOAD-NF-228-229", "invoice_ids": [invoice_id, second_invoice_id], "vehicle_plate": "abc1d23", "driver_name": "João Motorista", "observations": "Entregar pela doca 2"})
        self.assertEqual(created.status_code, 201, created.get_json())
        load_id = created.get_json()["id"]
        self.assertEqual(created.get_json()["items"], 1)
        self.assertEqual(created.get_json()["invoices"], ["228", "229"])
        self.assertEqual(created.get_json()["vehicle_plate"], "ABC1D23")
        before = self.client.get(f"/api/loads/{load_id}").get_json()
        self.assertEqual(before["invoices"][0]["number"], "228")
        self.assertEqual(before["vehicle_plate"], "ABC1D23")
        self.assertEqual(before["driver_name"], "João Motorista")
        self.assertEqual(before["observations"], "Entregar pela doca 2")
        self.assertIsNone(before["items"][0]["checked"])
        self.assertEqual(before["items"][0]["expected"], 6)
        stock_before = self.client.get("/api/products").get_json()[0]["stock"]
        pending_close = self.post_json(f"/api/loads/{load_id}/close", {})
        self.assertEqual(pending_close.status_code, 400)
        scanned = self.post_json(f"/api/loads/{load_id}/scan", {"barcode": "7891234567890", "quantity": 3})
        self.assertEqual(scanned.status_code, 201, scanned.get_json())
        self.assertEqual(scanned.get_json()["expected"], 6)
        detail = self.client.get(f"/api/loads/{load_id}").get_json()
        self.assertEqual(detail["items"][0]["checked"], 3)
        self.assertEqual(detail["items"][0]["expected"], 6)
        self.assertEqual(self.client.get("/api/products").get_json()[0]["stock"], stock_before)
        unconfirmed_close = self.post_json(f"/api/loads/{load_id}/close", {})
        self.assertEqual(unconfirmed_close.status_code, 409)
        self.assertTrue(unconfirmed_close.get_json()["requires_confirmation"])
        self.assertEqual(unconfirmed_close.get_json()["summary"]["discrepancies"], 1)
        confirmed_close = self.post_json(f"/api/loads/{load_id}/close", {"confirm_divergences": True})
        self.assertEqual(confirmed_close.status_code, 200)
        self.assertEqual(confirmed_close.get_json()["summary"]["expected"], 6)

    def test_conference_quantity_can_be_corrected_and_last_item_undone_with_audit(self):
        created = self.post_json("/api/loads", {"reference": "LOAD-EDIT", "supplier": "Fornecedor Alfa"})
        load_id = created.get_json()["id"]
        scanned = self.post_json(f"/api/loads/{load_id}/scan", {"barcode": "7891234567890", "expected_quantity": 8, "quantity": 7})
        self.assertEqual(scanned.status_code, 201)
        item_id = scanned.get_json()["id"]
        edited = self.client.put(f"/api/loads/{load_id}/items/{item_id}", json={"quantity": 6}, headers={"X-CSRFToken": self.csrf})
        self.assertEqual(edited.status_code, 200)
        self.assertEqual(edited.get_json()["checked"], 6)
        undone = self.post_json(f"/api/loads/{load_id}/undo-last", {})
        self.assertEqual(undone.status_code, 200)
        self.assertEqual(undone.get_json()["quantity"], 6)
        detail = self.client.get(f"/api/loads/{load_id}").get_json()
        self.assertIsNone(detail["items"][0]["checked"])
        with application.app.app_context():
            actions = [entry.action for entry in application.AuditLog.query.filter_by(entity="carga").all()]
        self.assertIn("editar_conferencia", actions)
        self.assertIn("desfazer_conferencia", actions)
        self.assertEqual(self.post_json(f"/api/loads/{load_id}/scan", {"barcode": "7891234567890", "expected_quantity": 8, "quantity": 7}).status_code, 201)
        zero = self.client.put(f"/api/loads/{load_id}/items/{item_id}", json={"quantity": 0}, headers={"X-CSRFToken": self.csrf})
        self.assertEqual(zero.status_code, 200)
        zero_detail = self.client.get(f"/api/loads/{load_id}").get_json()
        self.assertEqual(zero_detail["items"][0]["checked"], 0)
        close = self.post_json(f"/api/loads/{load_id}/close", {})
        self.assertEqual(close.status_code, 409)
        self.assertEqual(close.get_json()["summary"]["discrepancies"], 1)

    def test_invoice_xml_import_is_idempotent_and_updates_stock(self):
        xml = b'''<NFe><infNFe Id="NFe12345678901234567890123456789012345678901234"><ide><nNF>220</nNF></ide><emit><xNome>Fornecedor XML</xNome></emit><det><prod><cProd>PRD-001</cProd><cEAN>7891234567890</cEAN><xProd>Leitor industrial</xProd><qCom>4.0000</qCom><vUnCom>10.25</vUnCom></prod></det></infNFe></NFe>'''
        response = self.client.post("/api/invoices/import", data={"file": (BytesIO(xml), "nf.xml")}, headers={"X-CSRFToken": self.csrf}, content_type="multipart/form-data")
        self.assertEqual(response.status_code, 201, response.get_json())
        duplicate = self.client.post("/api/invoices/import", data={"file": (BytesIO(xml), "nf.xml")}, headers={"X-CSRFToken": self.csrf}, content_type="multipart/form-data")
        self.assertEqual(duplicate.status_code, 400)
        product = self.client.get("/api/products").get_json()[0]
        self.assertEqual(product["stock"], 19)

    def test_multiple_invoice_xml_files_import_in_one_request(self):
        first_xml = b'''<NFe><infNFe Id="NFe22345678901234567890123456789012345678901234"><ide><nNF>221</nNF></ide><emit><xNome>Fornecedor XML</xNome></emit><det><prod><cProd>PRD-001</cProd><cEAN>7891234567890</cEAN><xProd>Leitor industrial</xProd><qCom>4</qCom><vUnCom>10</vUnCom></prod></det></infNFe></NFe>'''
        second_xml = b'''<NFe><infNFe Id="NFe32345678901234567890123456789012345678901234"><ide><nNF>222</nNF></ide><emit><xNome>Fornecedor XML</xNome></emit><det><prod><cProd>PRD-001</cProd><cEAN>7891234567890</cEAN><xProd>Leitor industrial</xProd><qCom>2</qCom><vUnCom>10</vUnCom></prod></det></infNFe></NFe>'''
        response = self.client.post("/api/invoices/import", data={"files": [(BytesIO(first_xml), "nf-221.xml"), (BytesIO(second_xml), "nf-222.xml")]}, headers={"X-CSRFToken": self.csrf}, content_type="multipart/form-data")
        self.assertEqual(response.status_code, 200, response.get_json())
        result = response.get_json()
        self.assertEqual(result["imported"], 2)
        self.assertEqual(result["failed"], 0)
        self.assertEqual(self.client.get("/api/invoices").get_json().__len__(), 2)
        self.assertEqual(self.client.get("/api/products").get_json()[0]["stock"], 21)

    def test_invoice_creates_products_supplier_and_origin_reference_automatically(self):
        xml = b'''<NFe><infNFe Id="NFe62345678901234567890123456789012345678901234"><ide><nNF>225</nNF></ide><emit><CNPJ>12.345.678/0001-90</CNPJ><xNome>Fornecedor Novo</xNome></emit><det><prod><cProd>NOVO-001</cProd><cEAN>7901234567890</cEAN><xProd>Mouse sem fio</xProd><uCom>cx</uCom><qCom>6</qCom><vUnCom>25</vUnCom></prod></det></infNFe></NFe>'''
        response = self.client.post("/api/invoices/import", data={"file": (BytesIO(xml), "nf-225.xml")}, headers={"X-CSRFToken": self.csrf}, content_type="multipart/form-data")
        self.assertEqual(response.status_code, 201, response.get_json())
        result = self.client.get("/api/products?q=NOVO-001").get_json()
        self.assertEqual(len(result), 1)
        product = result[0]
        self.assertEqual(product["name"], "Mouse sem fio")
        self.assertEqual(product["unit"], "cx")
        self.assertEqual(product["stock"], 6)
        self.assertEqual(product["supplier"], "Fornecedor Novo")
        self.assertEqual(product["source_invoice_number"], "225")
        suppliers = self.client.get("/api/suppliers").get_json()
        supplier = next(item for item in suppliers if item["name"] == "Fornecedor Novo")
        self.assertEqual(supplier["document"], "12345678000190")

    def test_same_supplier_code_from_different_suppliers_creates_distinct_products(self):
        first_xml = b'''<NFe><infNFe Id="NFe72345678901234567890123456789012345678901234"><ide><nNF>226</nNF></ide><emit><CNPJ>11111111000111</CNPJ><xNome>Fornecedor Um</xNome></emit><det><prod><cProd>EXT-01</cProd><cEAN>7900000000011</cEAN><xProd>Produto do fornecedor um</xProd><uCom>un</uCom><qCom>2</qCom><vUnCom>10</vUnCom></prod></det></infNFe></NFe>'''
        second_xml = b'''<NFe><infNFe Id="NFe82345678901234567890123456789012345678901234"><ide><nNF>227</nNF></ide><emit><CNPJ>22222222000122</CNPJ><xNome>Fornecedor Dois</xNome></emit><det><prod><cProd>EXT-01</cProd><cEAN>7900000000028</cEAN><xProd>Produto do fornecedor dois</xProd><uCom>un</uCom><qCom>3</qCom><vUnCom>10</vUnCom></prod></det></infNFe></NFe>'''
        response = self.client.post("/api/invoices/import", data={"files": [(BytesIO(first_xml), "nf-226.xml"), (BytesIO(second_xml), "nf-227.xml")]}, headers={"X-CSRFToken": self.csrf}, content_type="multipart/form-data")
        self.assertEqual(response.status_code, 200, response.get_json())
        first = self.client.get("/api/products?q=Produto%20do%20fornecedor%20um").get_json()
        second = self.client.get("/api/products?q=Produto%20do%20fornecedor%20dois").get_json()
        self.assertEqual(len(first), 1)
        self.assertEqual(len(second), 1)
        self.assertNotEqual(first[0]["code"], second[0]["code"])

    def test_delete_invoice_cancels_and_reverses_only_its_stock_entries(self):
        xml = b'''<NFe><infNFe Id="NFe42345678901234567890123456789012345678901234"><ide><nNF>223</nNF></ide><emit><xNome>Fornecedor XML</xNome></emit><det><prod><cProd>PRD-001</cProd><cEAN>7891234567890</cEAN><xProd>Leitor industrial</xProd><qCom>4</qCom><vUnCom>10</vUnCom></prod></det></infNFe></NFe>'''
        imported = self.client.post("/api/invoices/import", data={"file": (BytesIO(xml), "nf.xml")}, headers={"X-CSRFToken": self.csrf}, content_type="multipart/form-data")
        invoice_id = imported.get_json()["id"]
        missing_token = self.client.delete(f"/api/invoices/{invoice_id}")
        self.assertEqual(missing_token.status_code, 400)
        deleted = self.client.delete(f"/api/invoices/{invoice_id}", headers={"X-CSRFToken": self.csrf})
        self.assertEqual(deleted.status_code, 200, deleted.get_json())
        self.assertEqual(self.client.get("/api/products").get_json()[0]["stock"], 15)
        self.assertEqual(self.client.get("/api/invoices").get_json(), [])
        duplicate = self.client.post("/api/invoices/import", data={"file": (BytesIO(xml), "nf.xml")}, headers={"X-CSRFToken": self.csrf}, content_type="multipart/form-data")
        self.assertEqual(duplicate.status_code, 400)

    def test_invoice_delete_is_blocked_when_receipt_stock_was_consumed(self):
        xml = b'''<NFe><infNFe Id="NFe52345678901234567890123456789012345678901234"><ide><nNF>224</nNF></ide><emit><xNome>Fornecedor XML</xNome></emit><det><prod><cProd>PRD-001</cProd><cEAN>7891234567890</cEAN><xProd>Leitor industrial</xProd><qCom>10</qCom><vUnCom>10</vUnCom></prod></det></infNFe></NFe>'''
        imported = self.client.post("/api/invoices/import", data={"file": (BytesIO(xml), "nf.xml")}, headers={"X-CSRFToken": self.csrf}, content_type="multipart/form-data")
        invoice_id = imported.get_json()["id"]
        self.post_json("/api/movements", {"product_id": self.product_id, "kind": "saida", "quantity": 16})
        deleted = self.client.delete(f"/api/invoices/{invoice_id}", headers={"X-CSRFToken": self.csrf})
        self.assertEqual(deleted.status_code, 409)
        self.assertEqual(self.client.get("/api/products").get_json()[0]["stock"], 9)

    def test_inventory_counts_and_adjusts_stock_when_closed(self):
        created = self.post_json("/api/inventories", {"reference": "INV-001"})
        inventory_id = created.get_json()["id"]
        detail = self.client.get(f"/api/inventories/{inventory_id}").get_json()
        item = detail["items"][0]
        counted = self.client.put(f"/api/inventories/{inventory_id}/items/{item['id']}", json={"counted": 12}, headers={"X-CSRFToken": self.csrf})
        self.assertEqual(counted.status_code, 200)
        closed = self.post_json(f"/api/inventories/{inventory_id}/close", {})
        self.assertEqual(closed.status_code, 200)
        product = self.client.get("/api/products").get_json()[0]
        self.assertEqual(product["stock"], 12)

    def test_reports_filter_and_export_csv_excel(self):
        self.post_json("/api/movements", {"product_id": self.product_id, "kind": "entrada", "quantity": 5, "reference": "COMPRA"})
        filters = f"?product_id={self.product_id}&category_id={self.category_id}&supplier_id={self.supplier_id}&responsible_id={self.admin_id}"
        report = self.client.get(f"/api/reports/movements{filters}")
        self.assertEqual(report.status_code, 200)
        self.assertEqual(len(report.get_json()), 2)
        csv_response = self.client.get(f"/api/reports/export/csv{filters}")
        self.assertEqual(csv_response.status_code, 200)
        self.assertIn("text/csv", csv_response.content_type)
        excel_response = self.client.get(f"/api/reports/export/xlsx{filters}")
        self.assertEqual(excel_response.status_code, 200)
        self.assertTrue(excel_response.data.startswith(b"PK"))
        workbook = load_workbook(BytesIO(excel_response.data), read_only=True)
        self.assertEqual(workbook.active.max_row, 3)
        today = datetime.now(timezone.utc).date().isoformat()
        self.assertEqual(len(self.client.get(f"/api/reports/movements?from={today}&to={today}").get_json()), 2)
        self.assertEqual(self.client.get("/api/reports/movements?from=2030-02-01&to=2030-01-01").status_code, 400)

    def test_exports_neutralize_spreadsheet_formulas(self):
        with application.app.app_context():
            product = db.session.get(Product, self.product_id)
            product.name = "=SUM(1,1)"
            db.session.commit()
        csv_response = self.client.get("/api/reports/export/csv")
        self.assertIn(b"'=SUM(1,1)", csv_response.data)
        excel_response = self.client.get("/api/reports/export/xlsx")
        workbook = load_workbook(BytesIO(excel_response.data), read_only=True)
        self.assertEqual(workbook.active["B2"].value, "'=SUM(1,1)")

    def test_user_updates_are_audited_and_deactivation_revokes_access(self):
        client = application.app.test_client()
        login_page = client.get("/login")
        token = re.search(rb'name="csrf_token" value="([^"]+)"', login_page.data).group(1).decode()
        client.post("/login", data={"csrf_token": token, "email": "caio@example.com", "password": "SenhaSegura123"})
        admin_page = self.client.get("/")
        admin_token = re.search(rb'name="csrf-token" content="([^"]+)"', admin_page.data).group(1).decode()
        updated = self.client.put("/api/users/2", json={"name": "Caio Conferente", "role": "conferente", "active": False}, headers={"X-CSRFToken": admin_token})
        self.assertEqual(updated.status_code, 200)
        self.assertEqual(client.get("/api/me").status_code, 302)
        last_admin = self.client.put(f"/api/users/{self.admin_id}", json={"name": "Ana Administradora", "role": "conferente", "active": True}, headers={"X-CSRFToken": admin_token})
        self.assertEqual(last_admin.status_code, 400)
        audit = self.client.get("/api/audit").get_json()
        self.assertTrue(any(entry["action"] == "editar" and entry["entity"] == "usuario" for entry in audit))

    def test_backup_can_be_downloaded_and_restored(self):
        backup = self.client.get("/api/settings/backup")
        self.assertEqual(backup.status_code, 200)
        self.assertTrue(backup.data.startswith(b"SQLite format 3"))
        backup_bytes = backup.data
        backup.close()
        restored = self.client.post("/api/settings/restore", data={"file": (BytesIO(backup_bytes), "snapshot.sqlite")}, headers={"X-CSRFToken": self.csrf}, content_type="multipart/form-data")
        self.assertEqual(restored.status_code, 200, restored.get_json())
        self.assertEqual(self.client.get("/api/products").status_code, 200)

    def test_conferente_can_read_but_cannot_mutate_stock(self):
        client = application.app.test_client()
        login_page = client.get("/login")
        token = re.search(rb'name="csrf_token" value="([^"]+)"', login_page.data).group(1).decode()
        response = client.post("/login", data={"csrf_token": token, "email": "caio@example.com", "password": "SenhaSegura123"})
        self.assertEqual(response.status_code, 302)
        page = client.get("/")
        token = re.search(rb'name="csrf-token" content="([^"]+)"', page.data).group(1).decode()
        self.assertEqual(client.get("/api/products").status_code, 200)
        denied = client.post("/api/movements", json={"product_id": self.product_id, "kind": "saida", "quantity": 1}, headers={"X-CSRFToken": token})
        self.assertEqual(denied.status_code, 403)

    def test_session_survives_page_refresh_and_logout_revokes_access(self):
        self.assertEqual(self.client.get("/api/me").status_code, 200)
        self.assertEqual(self.client.get("/").status_code, 200)
        self.client.post("/logout", data={"csrf_token": self.csrf})
        self.assertEqual(self.client.get("/api/me").status_code, 302)
        self.client.post("/login", data={"csrf_token": self.csrf, "email": "ana@example.com", "password": "SenhaSegura123"})
        self.assertEqual(self.client.get("/api/products").get_json()[0]["stock"], 15)


if __name__ == "__main__":
    unittest.main()