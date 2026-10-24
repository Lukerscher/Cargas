import csv
import io
import json
import math
import os
import secrets
import sqlite3
from collections import Counter
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import urlparse
from xml.etree import ElementTree as StdET

from defusedxml import ElementTree as SafeET
from flask import Flask, jsonify, make_response, redirect, render_template, request, send_file, url_for
from flask_login import LoginManager, UserMixin, current_user, login_required, login_user, logout_user
from flask_sqlalchemy import SQLAlchemy
from flask_wtf.csrf import CSRFError, CSRFProtect, generate_csrf
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import event, func, inspect, or_
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename


BASE_DIR = Path(__file__).resolve().parent
INSTANCE_DIR = BASE_DIR / "instance"
INSTANCE_DIR.mkdir(exist_ok=True)
DB_PATH = Path(os.environ.get("DATABASE_PATH", INSTANCE_DIR / "almoxcontrol.sqlite")).resolve()
DB_PATH.parent.mkdir(parents=True, exist_ok=True)
SECRET_PATH = INSTANCE_DIR / ".secret_key"
if not os.environ.get("SECRET_KEY") and not SECRET_PATH.exists():
    SECRET_PATH.write_text(secrets.token_hex(32), encoding="ascii")

app = Flask(__name__, instance_path=str(INSTANCE_DIR))
app.config.update(
    SECRET_KEY=os.environ.get("SECRET_KEY") or SECRET_PATH.read_text(encoding="ascii"),
    SQLALCHEMY_DATABASE_URI=f"sqlite:///{DB_PATH}",
    SQLALCHEMY_TRACK_MODIFICATIONS=False,
    MAX_CONTENT_LENGTH=250 * 1024 * 1024,
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=os.environ.get("FLASK_ENV") == "production",
    WTF_CSRF_TIME_LIMIT=3600,
)
db = SQLAlchemy(app)
csrf = CSRFProtect(app)
login_manager = LoginManager(app)
login_manager.login_view = "login"
login_manager.login_message = "Entre para acessar o sistema."
login_manager.login_message_category = "info"

ROLES = {"administrador", "almoxarife", "conferente"}
ROLE_LABELS = {"administrador": "Administrador", "almoxarife": "Almoxarife", "conferente": "Conferente"}


def now_utc():
    return datetime.now(timezone.utc).replace(tzinfo=None)


class User(UserMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(120), nullable=False)
    email = db.Column(db.String(254), unique=True, nullable=False, index=True)
    password_hash = db.Column(db.String(256), nullable=False)
    role = db.Column(db.String(24), nullable=False, default="conferente")
    active = db.Column(db.Boolean, nullable=False, default=True)
    created_at = db.Column(db.DateTime, nullable=False, default=now_utc)

    def set_password(self, password):
        if len(password) < 10:
            raise ValueError("A senha deve ter pelo menos 10 caracteres.")
        self.password_hash = generate_password_hash(password, method="scrypt")

    def public(self):
        return {"id": self.id, "name": self.name, "email": self.email, "role": self.role, "role_label": ROLE_LABELS[self.role], "active": self.active}


class AuditLog(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    actor_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False, index=True)
    action = db.Column(db.String(40), nullable=False)
    entity = db.Column(db.String(40), nullable=False, index=True)
    record_id = db.Column(db.String(80), nullable=False)
    details = db.Column(db.Text, nullable=False, default="{}")
    created_at = db.Column(db.DateTime, nullable=False, default=now_utc, index=True)
    actor = db.relationship(User)


class Category(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(100), unique=True, nullable=False)


class Supplier(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(160), nullable=False, index=True)
    document = db.Column(db.String(32), unique=True)
    email = db.Column(db.String(254))
    phone = db.Column(db.String(32))
    active = db.Column(db.Boolean, nullable=False, default=True)


class Product(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    code = db.Column(db.String(80), unique=True, nullable=False, index=True)
    barcode = db.Column(db.String(80), unique=True, index=True)
    name = db.Column(db.String(180), nullable=False, index=True)
    unit = db.Column(db.String(12), nullable=False, default="un")
    min_stock = db.Column(db.Float, nullable=False, default=0)
    active = db.Column(db.Boolean, nullable=False, default=True)
    category_id = db.Column(db.Integer, db.ForeignKey("category.id"))
    supplier_id = db.Column(db.Integer, db.ForeignKey("supplier.id"))
    source_invoice_id = db.Column(db.Integer, db.ForeignKey("invoice.id"), index=True)
    category = db.relationship(Category)
    supplier = db.relationship(Supplier)
    source_invoice = db.relationship("Invoice", foreign_keys=[source_invoice_id])
    created_at = db.Column(db.DateTime, nullable=False, default=now_utc)


class Movement(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    product_id = db.Column(db.Integer, db.ForeignKey("product.id"), nullable=False, index=True)
    kind = db.Column(db.String(12), nullable=False)
    quantity = db.Column(db.Float, nullable=False)
    reference = db.Column(db.String(180))
    responsible_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False, index=True)
    invoice_id = db.Column(db.Integer, db.ForeignKey("invoice.id"), index=True)
    created_at = db.Column(db.DateTime, nullable=False, default=now_utc, index=True)
    product = db.relationship(Product)
    responsible = db.relationship(User)
    invoice = db.relationship("Invoice", foreign_keys=[invoice_id])


class Invoice(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    access_key = db.Column(db.String(60), unique=True, nullable=False, index=True)
    number = db.Column(db.String(40), nullable=False)
    supplier_name = db.Column(db.String(180), nullable=False)
    issued_at = db.Column(db.String(30))
    filename = db.Column(db.String(180), nullable=False)
    imported_by_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    imported_at = db.Column(db.DateTime, nullable=False, default=now_utc)
    status = db.Column(db.String(16), nullable=False, default="importada")
    cancelled_at = db.Column(db.DateTime)
    cancelled_by_id = db.Column(db.Integer, db.ForeignKey("user.id"))
    imported_by = db.relationship(User, foreign_keys=[imported_by_id])
    cancelled_by = db.relationship(User, foreign_keys=[cancelled_by_id])
    items = db.relationship("InvoiceItem", cascade="all, delete-orphan", back_populates="invoice")


class InvoiceItem(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    invoice_id = db.Column(db.Integer, db.ForeignKey("invoice.id"), nullable=False)
    product_id = db.Column(db.Integer, db.ForeignKey("product.id"))
    supplier_code = db.Column(db.String(80))
    description = db.Column(db.String(180), nullable=False)
    quantity = db.Column(db.Float, nullable=False)
    unit_value = db.Column(db.Float, nullable=False, default=0)
    invoice = db.relationship(Invoice, back_populates="items")
    product = db.relationship(Product)


class Load(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    reference = db.Column(db.String(80), unique=True, nullable=False)
    supplier = db.Column(db.String(180), nullable=False)
    status = db.Column(db.String(16), nullable=False, default="aberta")
    created_by_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    created_at = db.Column(db.DateTime, nullable=False, default=now_utc)
    created_by = db.relationship(User)
    items = db.relationship("ConferenceItem", cascade="all, delete-orphan", back_populates="load")
    invoice_links = db.relationship("LoadInvoice", cascade="all, delete-orphan", back_populates="load")


class LoadInvoice(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    load_id = db.Column(db.Integer, db.ForeignKey("load.id", ondelete="CASCADE"), nullable=False)
    invoice_id = db.Column(db.Integer, db.ForeignKey("invoice.id"), nullable=False)
    load = db.relationship(Load, back_populates="invoice_links")
    invoice = db.relationship(Invoice)
    __table_args__ = (db.UniqueConstraint("load_id", "invoice_id", name="uq_load_invoice"),)


class ConferenceItem(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    load_id = db.Column(db.Integer, db.ForeignKey("load.id"), nullable=False)
    product_id = db.Column(db.Integer, db.ForeignKey("product.id"), nullable=False)
    expected_quantity = db.Column(db.Float, nullable=False, default=0)
    checked_quantity = db.Column(db.Float, nullable=False)
    checked_by_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    checked_at = db.Column(db.DateTime, nullable=False, default=now_utc)
    load = db.relationship(Load, back_populates="items")
    product = db.relationship(Product)
    checked_by = db.relationship(User)
    __table_args__ = (db.UniqueConstraint("load_id", "product_id", name="uq_load_product"),)


class Inventory(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    reference = db.Column(db.String(80), unique=True, nullable=False)
    status = db.Column(db.String(16), nullable=False, default="aberto")
    created_by_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    created_at = db.Column(db.DateTime, nullable=False, default=now_utc)
    created_by = db.relationship(User)
    items = db.relationship("InventoryItem", cascade="all, delete-orphan", back_populates="inventory")


class InventoryItem(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    inventory_id = db.Column(db.Integer, db.ForeignKey("inventory.id"), nullable=False)
    product_id = db.Column(db.Integer, db.ForeignKey("product.id"), nullable=False)
    expected_quantity = db.Column(db.Float, nullable=False)
    counted_quantity = db.Column(db.Float)
    counted_by_id = db.Column(db.Integer, db.ForeignKey("user.id"))
    inventory = db.relationship(Inventory, back_populates="items")
    product = db.relationship(Product)
    counted_by = db.relationship(User)
    __table_args__ = (db.UniqueConstraint("inventory_id", "product_id", name="uq_inventory_product"),)


@login_manager.user_loader
def load_user(user_id):
    user = db.session.get(User, int(user_id))
    return user if user and user.active else None


def roles_required(*roles):
    def decorate(view):
        @wraps(view)
        @login_required
        def wrapped(*args, **kwargs):
            if current_user.role not in roles:
                return jsonify(error="Seu perfil não tem permissão para esta operação."), 403
            return view(*args, **kwargs)
        return wrapped
    return decorate


from functools import wraps


def payload():
    return request.get_json(silent=True) or {}


def audit(action, entity, record_id, details=None, actor=None):
    db.session.add(AuditLog(actor_id=(actor or current_user).id, action=action, entity=entity, record_id=str(record_id), details=json.dumps(details or {}, ensure_ascii=False, sort_keys=True)))


def required_text(data, key, label, maximum=180):
    value = str(data.get(key, "")).strip()
    if not value or len(value) > maximum:
        raise ValueError(f"Informe {label} (até {maximum} caracteres).")
    return value


def positive_number(value, label, allow_zero=False):
    try:
        number = float(value)
    except (TypeError, ValueError):
        raise ValueError(f"Informe um valor válido para {label}.")
    if not math.isfinite(number) or number < 0 or (number == 0 and not allow_zero):
        raise ValueError(f"{label} deve ser maior que zero.")
    return number


def stock_level(product_id):
    total = db.session.query(func.coalesce(func.sum(db.case(
        (Movement.kind == "entrada", Movement.quantity),
        (Movement.kind == "saida", -Movement.quantity),
        else_=Movement.quantity,
    )), 0)).filter(Movement.product_id == product_id).scalar()
    return float(total or 0)


def product_json(product):
    stock = stock_level(product.id)
    return {"id": product.id, "code": product.code, "barcode": product.barcode or "", "name": product.name, "unit": product.unit, "min_stock": product.min_stock, "stock": stock, "low_stock": stock <= product.min_stock, "active": product.active, "category_id": product.category_id, "category": product.category.name if product.category else "Sem categoria", "supplier_id": product.supplier_id, "supplier": product.supplier.name if product.supplier else "Sem fornecedor", "source_invoice_number": product.source_invoice.number if product.source_invoice else ""}


def filtered_movements(args):
    query = Movement.query.join(Product).join(User, Movement.responsible)
    from_date = date.fromisoformat(args["from"]) if args.get("from") else None
    to_date = date.fromisoformat(args["to"]) if args.get("to") else None
    if from_date and to_date and from_date > to_date:
        raise ValueError("A data inicial deve ser anterior à data final.")
    if from_date:
        query = query.filter(Movement.created_at >= datetime.combine(from_date, datetime.min.time()))
    if to_date:
        query = query.filter(Movement.created_at < datetime.combine(to_date + timedelta(days=1), datetime.min.time()))
    for key, column in (("product_id", Movement.product_id), ("category_id", Product.category_id), ("supplier_id", Product.supplier_id), ("responsible_id", Movement.responsible_id)):
        if args.get(key):
            value = int(args[key])
            if value < 1:
                raise ValueError("Filtro inválido.")
            query = query.filter(column == value)
    return query.order_by(Movement.created_at.desc())


def report_movements_rows(args):
    movements = filtered_movements(args).limit(5000).all()
    return [{"date": m.created_at.isoformat(timespec="seconds"), "product": m.product.name, "code": m.product.code, "category": m.product.category.name if m.product.category else "Sem categoria", "supplier": m.product.supplier.name if m.product.supplier else "Sem fornecedor", "kind": m.kind, "quantity": m.quantity, "unit": m.product.unit, "reference": m.reference or "", "responsible": m.responsible.name} for m in movements]


def safe_export_value(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


@app.errorhandler(CSRFError)
def handle_csrf_error(_error):
    return jsonify(error="Sessão expirada. Atualize a página e tente novamente."), 400


@app.errorhandler(413)
def file_too_large(_error):
    return jsonify(error="Arquivo excede o limite de 250 MB."), 413


@app.get("/login")
def login():
    if current_user.is_authenticated:
        return redirect(url_for("index"))
    return render_template("login.html", csrf_token=generate_csrf(), error=request.args.get("error"))


@app.post("/login")
def login_submit():
    email = request.form.get("email", "").strip().lower()
    user = User.query.filter_by(email=email, active=True).first()
    if not user or not check_password_hash(user.password_hash, request.form.get("password", "")):
        return redirect(url_for("login", error="E-mail ou senha inválidos."))
    audit("login", "usuario", user.id, actor=user)
    db.session.commit()
    login_user(user, remember=False, fresh=True)
    next_url = request.args.get("next", "")
    parsed_next = urlparse(next_url)
    if parsed_next.path.startswith("/") and not parsed_next.netloc and not parsed_next.scheme:
        return redirect(next_url)
    return redirect(url_for("index"))


@app.post("/logout")
@login_required
def logout():
    audit("logout", "usuario", current_user.id)
    db.session.commit()
    logout_user()
    return redirect(url_for("login"))


@app.get("/")
@login_required
def index():
    return render_template("app.html", csrf_token=generate_csrf(), user=current_user.public())


@app.get("/api/me")
@login_required
def api_me():
    return jsonify(current_user.public())


@app.get("/api/dashboard")
@login_required
def dashboard():
    active = Product.query.filter_by(active=True).order_by(Product.name).all()
    low = [product_json(p) for p in active if stock_level(p.id) <= p.min_stock]
    recent = Movement.query.order_by(Movement.created_at.desc()).limit(8).all()
    return jsonify({"products": len(active), "low_stock": len(low), "open_loads": Load.query.filter_by(status="aberta").count(), "open_inventories": Inventory.query.filter_by(status="aberto").count(), "low_products": low[:8], "recent": [{"id": m.id, "product": m.product.name, "kind": m.kind, "quantity": m.quantity, "unit": m.product.unit, "responsible": m.responsible.name, "created_at": m.created_at.isoformat(timespec="minutes")} for m in recent]})


@app.get("/api/products")
@login_required
def products_list():
    query = Product.query
    if request.args.get("active") != "todos":
        query = query.filter_by(active=True)
    term = request.args.get("q", "").strip()
    if term:
        pattern = f"%{term}%"
        query = query.filter(or_(Product.name.ilike(pattern), Product.code.ilike(pattern), Product.barcode.ilike(pattern)))
    return jsonify([product_json(p) for p in query.order_by(Product.name).all()])


@app.post("/api/products")
@roles_required("administrador", "almoxarife")
def product_create():
    data = payload()
    try:
        code = required_text(data, "code", "o código", 80)
        barcode = str(data.get("barcode", "")).strip() or None
        if Product.query.filter_by(code=code).first():
            raise ValueError("Já existe um produto com este código.")
        if barcode and Product.query.filter_by(barcode=barcode).first():
            raise ValueError("Este código de barras já está cadastrado.")
        category_id, supplier_id = data.get("category_id") or None, data.get("supplier_id") or None
        if category_id and not db.session.get(Category, int(category_id)):
            raise ValueError("A categoria informada não existe.")
        if supplier_id and not db.session.get(Supplier, int(supplier_id)):
            raise ValueError("O fornecedor informado não existe.")
        product = Product(code=code, barcode=barcode, name=required_text(data, "name", "o nome"), unit=required_text(data, "unit", "a unidade", 12), min_stock=positive_number(data.get("min_stock", 0), "Estoque mínimo", allow_zero=True), category_id=category_id, supplier_id=supplier_id)
        db.session.add(product)
        db.session.flush()
        audit("criar", "produto", product.id, {"codigo": code, "nome": product.name})
        db.session.commit()
        return jsonify(product_json(product)), 201
    except (ValueError, TypeError) as error:
        db.session.rollback()
        return jsonify(error=str(error)), 400


@app.put("/api/products/<int:product_id>")
@roles_required("administrador", "almoxarife")
def product_update(product_id):
    product, data = db.get_or_404(Product, product_id), payload()
    try:
        code, barcode = required_text(data, "code", "o código", 80), str(data.get("barcode", "")).strip() or None
        if Product.query.filter(Product.code == code, Product.id != product.id).first():
            raise ValueError("Já existe um produto com este código.")
        if barcode and Product.query.filter(Product.barcode == barcode, Product.id != product.id).first():
            raise ValueError("Este código de barras já está cadastrado.")
        category_id, supplier_id = data.get("category_id") or None, data.get("supplier_id") or None
        if category_id and not db.session.get(Category, int(category_id)):
            raise ValueError("A categoria informada não existe.")
        if supplier_id and not db.session.get(Supplier, int(supplier_id)):
            raise ValueError("O fornecedor informado não existe.")
        product.code, product.barcode = code, barcode
        product.name = required_text(data, "name", "o nome")
        product.unit = required_text(data, "unit", "a unidade", 12)
        product.min_stock = positive_number(data.get("min_stock", 0), "Estoque mínimo", allow_zero=True)
        product.category_id, product.supplier_id = category_id, supplier_id
        audit("editar", "produto", product.id, {"codigo": code, "nome": product.name})
        db.session.commit()
        return jsonify(product_json(product))
    except (ValueError, TypeError) as error:
        db.session.rollback()
        return jsonify(error=str(error)), 400


@app.delete("/api/products/<int:product_id>")
@roles_required("administrador", "almoxarife")
def product_archive(product_id):
    product = db.get_or_404(Product, product_id)
    product.active = False
    audit("arquivar", "produto", product.id, {"codigo": product.code, "nome": product.name})
    db.session.commit()
    return jsonify(ok=True)


@app.get("/api/categories")
@login_required
def categories_list():
    return jsonify([{"id": c.id, "name": c.name} for c in Category.query.order_by(Category.name).all()])


@app.post("/api/categories")
@roles_required("administrador", "almoxarife")
def category_create():
    try:
        name = required_text(payload(), "name", "a categoria", 100)
        if Category.query.filter(func.lower(Category.name) == name.lower()).first():
            raise ValueError("Esta categoria já está cadastrada.")
        category = Category(name=name)
        db.session.add(category)
        db.session.flush()
        audit("criar", "categoria", category.id, {"nome": name})
        db.session.commit()
        return jsonify({"id": category.id, "name": category.name}), 201
    except ValueError as error:
        db.session.rollback()
        return jsonify(error=str(error)), 400


@app.get("/api/suppliers")
@login_required
def suppliers_list():
    return jsonify([{"id": s.id, "name": s.name, "document": s.document or "", "email": s.email or "", "phone": s.phone or "", "active": s.active} for s in Supplier.query.order_by(Supplier.name).all()])


@app.post("/api/suppliers")
@roles_required("administrador", "almoxarife")
def supplier_create():
    data = payload()
    try:
        name, document = required_text(data, "name", "o fornecedor"), str(data.get("document", "")).strip() or None
        if document and Supplier.query.filter_by(document=document).first():
            raise ValueError("Este documento já pertence a um fornecedor cadastrado.")
        supplier = Supplier(name=name, document=document, email=str(data.get("email", "")).strip() or None, phone=str(data.get("phone", "")).strip() or None)
        db.session.add(supplier)
        db.session.flush()
        audit("criar", "fornecedor", supplier.id, {"nome": name})
        db.session.commit()
        return jsonify({"id": supplier.id, "name": name, "document": document or "", "email": supplier.email or "", "phone": supplier.phone or "", "active": True}), 201
    except ValueError as error:
        db.session.rollback()
        return jsonify(error=str(error)), 400


@app.get("/api/movements")
@login_required
def movements_list():
    try:
        movements = filtered_movements(request.args).limit(500).all()
    except (ValueError, TypeError):
        return jsonify(error="Filtro inválido."), 400
    return jsonify([{"id": m.id, "kind": m.kind, "quantity": m.quantity, "reference": m.reference or "", "product_id": m.product_id, "product": m.product.name, "code": m.product.code, "unit": m.product.unit, "responsible_id": m.responsible_id, "responsible": m.responsible.name, "created_at": m.created_at.isoformat(timespec="seconds")} for m in movements])


@app.post("/api/movements")
@roles_required("administrador", "almoxarife")
def movement_create():
    data = payload()
    product = db.session.get(Product, data.get("product_id"))
    kind = str(data.get("kind", ""))
    try:
        if not product or not product.active:
            raise ValueError("Selecione um produto ativo.")
        if kind not in {"entrada", "saida", "ajuste"}:
            raise ValueError("Tipo de movimentação inválido.")
        try:
            quantity = float(data.get("quantity", 0))
        except (ValueError, TypeError):
            raise ValueError("Informe uma quantidade válida.")
        if not math.isfinite(quantity) or not quantity or (kind != "ajuste" and quantity < 0):
            raise ValueError("A quantidade deve ser diferente de zero.")
        if kind == "saida" and stock_level(product.id) < quantity:
            raise ValueError("Estoque insuficiente para esta saída.")
        if kind == "ajuste" and stock_level(product.id) + quantity < 0:
            raise ValueError("O ajuste deixaria o estoque negativo.")
        movement = Movement(product=product, kind=kind, quantity=quantity, reference=str(data.get("reference", "")).strip()[:180] or None, responsible=current_user)
        db.session.add(movement)
        db.session.commit()
        return jsonify({"id": movement.id, "stock": stock_level(product.id)}), 201
    except (ValueError, TypeError) as error:
        db.session.rollback()
        return jsonify(error=str(error)), 400


def xml_text(element, tag):
    for child in element.iter():
        if child.tag.rsplit("}", 1)[-1] == tag and child.text:
            return child.text.strip()
    return ""


def supplier_from_invoice(root, name):
    emitter = next((node for node in root.iter() if node.tag.rsplit("}", 1)[-1] == "emit"), root)
    raw_document = xml_text(emitter, "CNPJ") or xml_text(emitter, "CPF")
    document = "".join(character for character in raw_document if character.isdigit()) or None
    supplier = Supplier.query.filter_by(document=document).first() if document else None
    if not supplier:
        supplier = Supplier.query.filter(func.lower(Supplier.name) == name.lower()).first()
    if not supplier:
        supplier = Supplier(name=name[:160], document=document)
        db.session.add(supplier)
        db.session.flush()
    return supplier


def product_from_invoice(code, barcode, description, unit, supplier, invoice):
    code = code.strip()[:80]
    barcode = barcode.strip()
    if barcode.upper() in {"", "SEM GTIN", "NAO INFORMADO", "N/A"} or set(barcode) <= {"0"}:
        barcode = ""
    if barcode:
        product = Product.query.filter_by(barcode=barcode).first()
        if product:
            if not product.supplier_id:
                product.supplier = supplier
            if not product.source_invoice_id:
                product.source_invoice = invoice
            return product, False
    product = None
    if code:
        if supplier.document:
            product_code = f"{supplier.id}-{code}"[:80]
            product = Product.query.filter_by(code=product_code).first()
        else:
            product_code = code
        if not product:
            product = Product.query.filter_by(code=code).first()
            if product and product.supplier_id not in (None, supplier.id):
                product = None
    else:
        product_code = f"EAN-{barcode}"[:80] if barcode else ""
    if product:
        if not product.supplier_id:
            product.supplier = supplier
        if not product.source_invoice_id:
            product.source_invoice = invoice
        if barcode and not product.barcode:
            product.barcode = barcode
        return product, False
    if not product_code:
        raise ValueError(f"o item '{description}' não possui cProd nem cEAN para cadastro automático")
    if Product.query.filter_by(code=product_code).first():
        product_code = f"{supplier.id}-{code}"[:80] if code else f"{supplier.id}-EAN-{barcode}"[:80]
    if not product_code or Product.query.filter_by(code=product_code).first():
        raise ValueError(f"não foi possível gerar um código único para '{description}'")
    product = Product(code=product_code, barcode=barcode or None, name=description[:180],
                      unit=(unit.strip() or "un")[:12], supplier=supplier, source_invoice=invoice)
    db.session.add(product)
    db.session.flush()
    audit("criar", "produto", product.id,
          {"codigo": product.code, "codigo_fornecedor": code, "nome": product.name, "nf": invoice.number, "fornecedor": supplier.name})
    return product, True


def import_invoice_file(uploaded):
    if not uploaded or not uploaded.filename or not uploaded.filename.lower().endswith(".xml"):
        raise ValueError("selecione um arquivo XML de NF-e")
    data = uploaded.stream.read(10 * 1024 * 1024 + 1)
    if len(data) > 10 * 1024 * 1024:
        raise ValueError("o XML excede o limite de 10 MB")
    root = SafeET.fromstring(data)
    info = next((node for node in root.iter() if node.tag.rsplit("}", 1)[-1] == "infNFe"), root)
    access_key = str(info.attrib.get("Id", "")).removeprefix("NFe") or xml_text(root, "chNFe")
    if len(access_key) < 20:
        raise ValueError("a chave de acesso da NF-e não foi encontrada")
    if Invoice.query.filter_by(access_key=access_key).first():
        raise ValueError("esta nota fiscal já foi importada")
    supplier_name = xml_text(root, "xNome") or "Fornecedor não identificado"
    supplier = supplier_from_invoice(root, supplier_name)
    number = xml_text(root, "nNF") or access_key[-9:]
    invoice = Invoice(access_key=access_key, number=number, supplier_name=supplier_name,
                      issued_at=xml_text(root, "dhEmi") or xml_text(root, "dEmi"),
                      filename=secure_filename(uploaded.filename)[:180] or "nota.xml", imported_by=current_user)
    db.session.add(invoice)
    db.session.flush()
    items = []
    for detail in (node for node in root.iter() if node.tag.rsplit("}", 1)[-1] == "det"):
        code, description = xml_text(detail, "cProd"), xml_text(detail, "xProd")
        quantity_text, value_text = xml_text(detail, "qCom"), xml_text(detail, "vUnCom")
        if not description or not quantity_text:
            continue
        quantity = positive_number(quantity_text, "Quantidade da nota")
        unit_value = positive_number(value_text or 0, "Valor unitário", allow_zero=True)
        barcode = xml_text(detail, "cEAN") or xml_text(detail, "cEANTrib")
        unit = xml_text(detail, "uCom") or xml_text(detail, "uTrib") or "un"
        product, created = product_from_invoice(code, barcode, description, unit, supplier, invoice)
        items.append((InvoiceItem(invoice=invoice, product=product, supplier_code=code[:80] or None,
                      description=description[:180], quantity=quantity, unit_value=unit_value), product, created))
    if not items:
        raise ValueError("nenhum item válido foi encontrado no XML")
    matched = 0
    for item, product, _created in items:
        db.session.add(item)
        if product and product.active:
            db.session.add(Movement(product=product, kind="entrada", quantity=item.quantity,
                                    reference=f"NF-e {number}", responsible=current_user, invoice=invoice))
            matched += 1
    audit("importar", "nota_fiscal", invoice.id, {"numero": number, "chave": access_key, "itens": len(items)})
    db.session.commit()
    return {"id": invoice.id, "number": number, "items": len(items), "stocked": matched, "unmatched": len(items) - matched}


@app.get("/api/invoices")
@login_required
def invoices_list():
    invoices = Invoice.query.filter_by(status="importada").order_by(Invoice.imported_at.desc()).limit(300).all()
    return jsonify([{"id": i.id, "number": i.number, "supplier": i.supplier_name, "issued_at": i.issued_at or "", "access_key": i.access_key, "filename": i.filename, "items": len(i.items), "imported_by": i.imported_by.name, "imported_at": i.imported_at.isoformat(timespec="minutes"), "status": i.status} for i in invoices])


@app.post("/api/invoices/import")
@roles_required("administrador", "almoxarife")
def invoice_import():
    uploads = request.files.getlist("files") or request.files.getlist("file")
    if not uploads:
        return jsonify(error="Selecione ao menos um arquivo XML de nota fiscal."), 400
    if len(uploads) > 20:
        return jsonify(error="Importe no máximo 20 notas por lote."), 400
    results = []
    for uploaded in uploads:
        try:
            result = import_invoice_file(uploaded)
            results.append({"filename": secure_filename(uploaded.filename), "ok": True, **result})
        except (ValueError, TypeError, StdET.ParseError) as error:
            db.session.rollback()
            results.append({"filename": secure_filename(uploaded.filename or "arquivo"), "ok": False, "error": str(error)})
    imported = sum(1 for result in results if result["ok"])
    failed = len(results) - imported
    response = {"results": results, "imported": imported, "failed": failed}
    if len(results) == 1 and imported:
        response.update({key: value for key, value in results[0].items() if key not in {"ok", "filename"}})
        return jsonify(response), 201
    return jsonify(response), 207 if imported and failed else 200 if imported else 400


@app.delete("/api/invoices/<int:invoice_id>")
@roles_required("administrador", "almoxarife")
def invoice_delete(invoice_id):
    invoice = db.get_or_404(Invoice, invoice_id)
    if invoice.status != "importada":
        return jsonify(error="Esta nota já foi excluída."), 400
    items = InvoiceItem.query.filter_by(invoice_id=invoice.id).filter(InvoiceItem.product_id.isnot(None)).all()
    incoming = Movement.query.filter_by(invoice_id=invoice.id, kind="entrada").all()
    expected_entries = Counter((item.product_id, item.quantity) for item in items if item.product)
    linked_entries = Counter((movement.product_id, movement.quantity) for movement in incoming)
    if expected_entries != linked_entries:
        return jsonify(error="Não foi possível identificar com segurança todas as entradas desta nota. Confira os movimentos vinculados antes de excluir."), 409
    totals = {}
    for movement in incoming:
        totals[movement.product_id] = totals.get(movement.product_id, 0) + movement.quantity
    for product_id, quantity in totals.items():
        if stock_level(product_id) < quantity:
            product = db.session.get(Product, product_id)
            return jsonify(error=f"Não é possível excluir a NF-e: o estoque de {product.name} já foi consumido. Faça um ajuste de estoque após conferir a operação."), 409
    for product_id, quantity in totals.items():
        db.session.add(Movement(product_id=product_id, kind="saida", quantity=quantity,
                                reference=f"Estorno NF-e {invoice.number} #{invoice.access_key[-8:]}",
                                responsible=current_user, invoice_id=invoice.id))
    invoice.status = "excluida"
    invoice.cancelled_at = now_utc()
    invoice.cancelled_by = current_user
    audit("excluir", "nota_fiscal", invoice.id, {"numero": invoice.number, "chave": invoice.access_key, "itens_estornados": len(totals)})
    db.session.commit()
    return jsonify(ok=True, message="Nota excluída e entradas de estoque estornadas.")


@app.get("/api/loads")
@login_required
def loads_list():
    loads = Load.query.order_by(Load.created_at.desc()).limit(300).all()
    return jsonify([{"id": load.id, "reference": load.reference, "supplier": load.supplier, "status": load.status, "created_by": load.created_by.name, "created_at": load.created_at.isoformat(timespec="minutes"), "items": len(load.items), "checked": sum(1 for item in load.items if item.checked_quantity > 0), "invoices": [link.invoice.number for link in load.invoice_links]} for load in loads])


@app.post("/api/loads")
@roles_required("administrador", "almoxarife", "conferente")
def load_create():
    data = payload()
    try:
        reference = required_text(data, "reference", "a referência", 80)
        if Load.query.filter_by(reference=reference).first():
            raise ValueError("Já existe uma carga com esta referência.")
        raw_invoice_ids = data.get("invoice_ids", [])
        if not isinstance(raw_invoice_ids, list):
            raise ValueError("Selecione notas fiscais válidas.")
        invoice_ids = list(dict.fromkeys(int(invoice_id) for invoice_id in raw_invoice_ids))
        invoices = Invoice.query.filter(Invoice.id.in_(invoice_ids), Invoice.status == "importada").all() if invoice_ids else []
        if len(invoices) != len(invoice_ids):
            raise ValueError("Uma ou mais notas selecionadas não estão disponíveis.")
        suppliers = sorted({invoice.supplier_name for invoice in invoices})
        supplier = " / ".join(suppliers) if suppliers else required_text(data, "supplier", "o fornecedor")
        load = Load(reference=reference, supplier=supplier[:180], created_by=current_user)
        db.session.add(load)
        db.session.flush()
        expected_by_product = {}
        for invoice in invoices:
            db.session.add(LoadInvoice(load=load, invoice=invoice))
            for invoice_item in invoice.items:
                product = invoice_item.product
                if product and product.active:
                    expected_by_product[product.id] = expected_by_product.get(product.id, 0) + invoice_item.quantity
        for product_id, expected_quantity in expected_by_product.items():
            db.session.add(ConferenceItem(load=load, product_id=product_id, expected_quantity=expected_quantity,
                                          checked_quantity=0, checked_by=current_user))
        audit("criar", "carga", load.id, {"referencia": reference, "notas": [invoice.number for invoice in invoices], "produtos_previstos": len(expected_by_product)})
        db.session.commit()
        return jsonify({"id": load.id, "reference": load.reference, "status": load.status,
                        "invoices": [invoice.number for invoice in invoices], "items": len(expected_by_product)}), 201
    except (ValueError, TypeError) as error:
        db.session.rollback()
        return jsonify(error=str(error)), 400


@app.get("/api/loads/<int:load_id>")
@login_required
def load_detail(load_id):
    load = db.get_or_404(Load, load_id)
    return jsonify({"id": load.id, "reference": load.reference, "supplier": load.supplier, "status": load.status,
                    "invoices": [{"id": link.invoice.id, "number": link.invoice.number, "supplier": link.invoice.supplier_name} for link in load.invoice_links],
                    "items": [{"id": item.id, "product_id": item.product_id, "code": item.product.code, "barcode": item.product.barcode or "", "name": item.product.name, "unit": item.product.unit, "expected": item.expected_quantity, "checked": item.checked_quantity if item.checked_quantity > 0 else None, "checked_by": item.checked_by.name if item.checked_quantity > 0 else ""} for item in load.items]})


@app.post("/api/loads/<int:load_id>/scan")
@roles_required("administrador", "almoxarife", "conferente")
def load_scan(load_id):
    load, data = db.get_or_404(Load, load_id), payload()
    if load.status != "aberta":
        return jsonify(error="Esta carga já foi encerrada."), 400
    product = Product.query.filter_by(barcode=str(data.get("barcode", "")).strip(), active=True).first()
    if not product:
        return jsonify(error="Código de barras não cadastrado em um produto ativo."), 404
    existing_item = ConferenceItem.query.filter_by(load_id=load.id, product_id=product.id).first()
    if existing_item and existing_item.checked_quantity > 0:
        return jsonify(error="Este produto já foi conferido nesta carga. A quantidade registrada não foi alterada."), 409
    try:
        quantity = positive_number(data.get("quantity"), "Quantidade conferida")
        expected = existing_item.expected_quantity if existing_item else positive_number(data.get("expected_quantity", quantity), "Quantidade esperada")
        if existing_item:
            existing_item.checked_quantity = quantity
            existing_item.checked_by = current_user
            existing_item.checked_at = now_utc()
            item = existing_item
        else:
            item = ConferenceItem(load=load, product=product, expected_quantity=expected, checked_quantity=quantity, checked_by=current_user)
            db.session.add(item)
        db.session.flush()
        audit("conferir", "carga", load.id, {"referencia": load.reference, "codigo": product.code, "previsto": expected, "recebido": quantity})
        db.session.commit()
        return jsonify({"id": item.id, "name": product.name, "code": product.code, "checked": quantity, "expected": expected}), 201
    except ValueError as error:
        db.session.rollback()
        return jsonify(error=str(error)), 400


@app.post("/api/loads/<int:load_id>/close")
@roles_required("administrador", "almoxarife")
def load_close(load_id):
    load = db.get_or_404(Load, load_id)
    if load.status != "aberta":
        return jsonify(error="A carga já foi encerrada."), 400
    pending = sum(1 for item in load.items if item.checked_quantity <= 0)
    if pending:
        return jsonify(error=f"Confira os {pending} produto(s) pendente(s) antes de encerrar a carga."), 400
    load.status = "conferida"
    audit("encerrar", "carga", load.id, {"referencia": load.reference})
    db.session.commit()
    return jsonify(ok=True)


@app.get("/api/inventories")
@login_required
def inventories_list():
    inventories = Inventory.query.order_by(Inventory.created_at.desc()).limit(300).all()
    return jsonify([{"id": inv.id, "reference": inv.reference, "status": inv.status, "created_by": inv.created_by.name, "created_at": inv.created_at.isoformat(timespec="minutes"), "items": len(inv.items), "counted": sum(1 for item in inv.items if item.counted_quantity is not None)} for inv in inventories])


@app.post("/api/inventories")
@roles_required("administrador", "almoxarife")
def inventory_create():
    try:
        reference = required_text(payload(), "reference", "a referência", 80)
        if Inventory.query.filter_by(reference=reference).first():
            raise ValueError("Já existe um inventário com esta referência.")
        inventory = Inventory(reference=reference, created_by=current_user)
        db.session.add(inventory)
        db.session.flush()
        for product in Product.query.filter_by(active=True).order_by(Product.name).all():
            db.session.add(InventoryItem(inventory=inventory, product=product, expected_quantity=stock_level(product.id)))
        audit("criar", "inventario", inventory.id, {"referencia": reference})
        db.session.commit()
        return jsonify({"id": inventory.id, "reference": reference, "items": len(inventory.items)}), 201
    except ValueError as error:
        db.session.rollback()
        return jsonify(error=str(error)), 400


@app.get("/api/inventories/<int:inventory_id>")
@login_required
def inventory_detail(inventory_id):
    inv = db.get_or_404(Inventory, inventory_id)
    return jsonify({"id": inv.id, "reference": inv.reference, "status": inv.status, "items": [{"id": item.id, "product_id": item.product_id, "code": item.product.code, "barcode": item.product.barcode or "", "name": item.product.name, "unit": item.product.unit, "expected": item.expected_quantity, "counted": item.counted_quantity} for item in inv.items]})


@app.put("/api/inventories/<int:inventory_id>/items/<int:item_id>")
@roles_required("administrador", "almoxarife")
def inventory_count(inventory_id, item_id):
    inv, item = db.get_or_404(Inventory, inventory_id), db.get_or_404(InventoryItem, item_id)
    if inv.id != item.inventory_id or inv.status != "aberto":
        return jsonify(error="Este item não pode ser contado."), 400
    try:
        item.counted_quantity = positive_number(payload().get("counted"), "Quantidade contada", allow_zero=True)
        item.counted_by = current_user
        audit("contar", "inventario", inv.id, {"referencia": inv.reference, "produto": item.product.code, "quantidade": item.counted_quantity})
        db.session.commit()
        return jsonify(ok=True)
    except ValueError as error:
        db.session.rollback()
        return jsonify(error=str(error)), 400


@app.post("/api/inventories/<int:inventory_id>/close")
@roles_required("administrador", "almoxarife")
def inventory_close(inventory_id):
    inventory = db.get_or_404(Inventory, inventory_id)
    if inventory.status != "aberto" or any(item.counted_quantity is None for item in inventory.items):
        return jsonify(error="Conte todos os produtos antes de fechar o inventário."), 400
    for item in inventory.items:
        difference = item.counted_quantity - item.expected_quantity
        if difference:
            db.session.add(Movement(product_id=item.product_id, kind="ajuste", quantity=difference, reference=f"Inventário {inventory.reference}", responsible=current_user))
    inventory.status = "concluido"
    audit("encerrar", "inventario", inventory.id, {"referencia": inventory.reference, "itens": len(inventory.items)})
    db.session.commit()
    return jsonify(ok=True)


@app.get("/api/reports/filters")
@login_required
def report_filters():
    return jsonify({"products": [{"id": p.id, "name": p.name} for p in Product.query.filter_by(active=True).order_by(Product.name).all()], "categories": [{"id": c.id, "name": c.name} for c in Category.query.order_by(Category.name).all()], "suppliers": [{"id": s.id, "name": s.name} for s in Supplier.query.order_by(Supplier.name).all()], "users": [u.public() for u in User.query.filter_by(active=True).order_by(User.name).all()]})


@app.get("/api/reports/movements")
@login_required
def report_movements():
    try:
        return jsonify(report_movements_rows(request.args))
    except (ValueError, TypeError):
        return jsonify(error="Filtro inválido."), 400


@app.get("/api/reports/export/<kind>")
@login_required
def report_export(kind):
    if kind not in {"csv", "xlsx"}:
        return jsonify(error="Formato não suportado."), 404
    try:
        rows = report_movements_rows(request.args)
    except (ValueError, TypeError):
        return jsonify(error="Filtro inválido."), 400
    headers = ["Data", "Produto", "Código", "Categoria", "Fornecedor", "Movimentação", "Quantidade", "Unidade", "Referência", "Responsável"]
    values = [[r["date"], r["product"], r["code"], r["category"], r["supplier"], r["kind"], r["quantity"], r["unit"], r["reference"], r["responsible"]] for r in rows]
    if kind == "csv":
        output = io.StringIO(newline="")
        output.write("\ufeff")
        writer = csv.writer(output, delimiter=";")
        writer.writerow(headers)
        writer.writerows([[safe_export_value(value) for value in row] for row in values])
        response = make_response(output.getvalue())
        response.headers["Content-Type"] = "text/csv; charset=utf-8"
        response.headers["Content-Disposition"] = "attachment; filename=relatorio-estoque.csv"
        return response
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Movimentações"
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E4D7A")
    for row in values:
        sheet.append([safe_export_value(value) for value in row])
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 38)
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name="relatorio-estoque.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/api/users")
@roles_required("administrador")
def users_list():
    return jsonify([u.public() for u in User.query.order_by(User.name).all()])


@app.post("/api/users")
@roles_required("administrador")
def user_create():
    data = payload()
    try:
        email = required_text(data, "email", "o e-mail", 254).lower()
        if User.query.filter_by(email=email).first():
            raise ValueError("Este e-mail já está cadastrado.")
        role = data.get("role", "conferente")
        if role not in ROLES:
            raise ValueError("Perfil de acesso inválido.")
        user = User(name=required_text(data, "name", "o nome"), email=email, role=role)
        user.set_password(str(data.get("password", "")))
        db.session.add(user)
        db.session.flush()
        audit("criar", "usuario", user.id, {"email": email, "perfil": role})
        db.session.commit()
        return jsonify(user.public()), 201
    except (ValueError, TypeError) as error:
        db.session.rollback()
        return jsonify(error=str(error)), 400


@app.put("/api/users/<int:user_id>")
@roles_required("administrador")
def user_update(user_id):
    user, data = db.get_or_404(User, user_id), payload()
    try:
        if user.id == current_user.id and not data.get("active", True):
            raise ValueError("Não é possível desativar seu próprio acesso.")
        was_active_admin = user.role == "administrador" and user.active
        role = data.get("role", user.role)
        if role not in ROLES:
            raise ValueError("Perfil de acesso inválido.")
        user.name, user.role = required_text(data, "name", "o nome"), role
        user.active = bool(data.get("active", user.active))
        if data.get("password"):
            user.set_password(str(data["password"]))
        if was_active_admin and (user.role != "administrador" or not user.active):
            remaining = User.query.filter_by(role="administrador", active=True).filter(User.id != user.id).count()
            if remaining == 0:
                raise ValueError("O sistema precisa manter ao menos um administrador ativo.")
        audit("editar", "usuario", user.id, {"email": user.email, "perfil": user.role, "ativo": user.active})
        db.session.commit()
        return jsonify(user.public())
    except ValueError as error:
        db.session.rollback()
        return jsonify(error=str(error)), 400


@app.get("/api/audit")
@roles_required("administrador")
def audit_list():
    entries = AuditLog.query.order_by(AuditLog.created_at.desc()).limit(300).all()
    return jsonify([{"id": e.id, "actor": e.actor.name, "action": e.action, "entity": e.entity, "record_id": e.record_id, "details": json.loads(e.details), "created_at": e.created_at.isoformat(timespec="seconds")} for e in entries])


@app.get("/api/settings/backup")
@roles_required("administrador")
def backup_download():
    backup_dir = INSTANCE_DIR / "backups"
    backup_dir.mkdir(exist_ok=True)
    path = backup_dir / f"almoxcontrol-{datetime.now().strftime('%Y%m%d-%H%M%S')}.sqlite"
    audit("backup", "banco", path.name)
    db.session.commit()
    with sqlite3.connect(DB_PATH) as source, sqlite3.connect(path) as destination:
        source.backup(destination)
    return send_file(path, as_attachment=True, download_name=path.name, mimetype="application/vnd.sqlite3")


@app.post("/api/settings/restore")
@roles_required("administrador")
def backup_restore():
    uploaded = request.files.get("file")
    if not uploaded or Path(uploaded.filename).suffix.lower() not in {".sqlite", ".db", ".backup"}:
        return jsonify(error="Selecione um backup SQLite válido (.sqlite, .db ou .backup)."), 400
    temporary = INSTANCE_DIR / f"restore-{secrets.token_hex(8)}.sqlite"
    try:
        uploaded.save(temporary)
        with sqlite3.connect(temporary) as source:
            check = source.execute("PRAGMA integrity_check").fetchone()[0]
            tables = {row[0] for row in source.execute("SELECT name FROM sqlite_master WHERE type='table'")}
            if check != "ok" or "user" not in tables or "product" not in tables or "audit_log" not in tables:
                raise ValueError("O arquivo não é um backup válido do AlmoxControl.")
            db.session.remove()
            db.engine.dispose()
            with sqlite3.connect(DB_PATH) as destination:
                source.backup(destination)
        return jsonify(ok=True, message="Backup restaurado. Entre novamente para carregar os dados recuperados.")
    except (sqlite3.Error, ValueError) as error:
        return jsonify(error=str(error)), 400
    finally:
        temporary.unlink(missing_ok=True)


@app.cli.command("create-admin")
def create_admin_command():
    from getpass import getpass
    email = input("E-mail do administrador: ").strip().lower()
    name = input("Nome: ").strip()
    password = getpass("Senha (mínimo 10 caracteres): ")
    with app.app_context():
        db.create_all()
        if User.query.filter_by(email=email).first():
            raise SystemExit("Este e-mail já está cadastrado.")
        user = User(name=name, email=email, role="administrador")
        user.set_password(password)
        db.session.add(user)
        db.session.flush()
        audit("criar", "usuario", user.id, {"email": email, "perfil": user.role}, actor=user)
        db.session.commit()
    print("Administrador criado.")


@app.cli.command("backup-db")
def backup_command():
    destination = INSTANCE_DIR / "backups" / f"almoxcontrol-{datetime.now().strftime('%Y%m%d-%H%M%S')}.sqlite"
    destination.parent.mkdir(exist_ok=True)
    with sqlite3.connect(DB_PATH) as source, sqlite3.connect(destination) as backup:
        source.backup(backup)
    print(destination)


@app.cli.command("restore-db")
def restore_command():
    backup_file = Path(input("Caminho do backup SQLite: ").strip()).expanduser().resolve()
    if not backup_file.is_file() or backup_file == DB_PATH:
        raise SystemExit("Informe um arquivo de backup válido e diferente do banco atual.")
    with sqlite3.connect(backup_file) as source:
        check = source.execute("PRAGMA integrity_check").fetchone()[0]
        if check != "ok":
            raise SystemExit("O backup falhou na verificação de integridade.")
        db.engine.dispose()
        with sqlite3.connect(DB_PATH) as destination:
            source.backup(destination)
    print("Backup restaurado. Reinicie a aplicação.")


with app.app_context():
    @event.listens_for(db.engine, "connect")
    def configure_sqlite(connection, _record):
        cursor = connection.cursor()
        cursor.execute("PRAGMA foreign_keys=ON")
        cursor.execute("PRAGMA journal_mode=WAL")
        cursor.execute("PRAGMA busy_timeout=5000")
        cursor.close()

    db.create_all()
    inspector = inspect(db.engine)
    movement_columns = {column["name"] for column in inspector.get_columns("movement")}
    invoice_columns = {column["name"] for column in inspector.get_columns("invoice")}
    product_columns = {column["name"] for column in inspector.get_columns("product")}
    with db.engine.begin() as connection:
        if "invoice_id" not in movement_columns:
            connection.exec_driver_sql("ALTER TABLE movement ADD COLUMN invoice_id INTEGER REFERENCES invoice(id)")
            connection.exec_driver_sql("CREATE INDEX IF NOT EXISTS ix_movement_invoice_id ON movement (invoice_id)")
        if "status" not in invoice_columns:
            connection.exec_driver_sql("ALTER TABLE invoice ADD COLUMN status VARCHAR(16) NOT NULL DEFAULT 'importada'")
        if "cancelled_at" not in invoice_columns:
            connection.exec_driver_sql("ALTER TABLE invoice ADD COLUMN cancelled_at DATETIME")
        if "cancelled_by_id" not in invoice_columns:
            connection.exec_driver_sql("ALTER TABLE invoice ADD COLUMN cancelled_by_id INTEGER REFERENCES user(id)")
        if "source_invoice_id" not in product_columns:
            connection.exec_driver_sql("ALTER TABLE product ADD COLUMN source_invoice_id INTEGER REFERENCES invoice(id)")
            connection.exec_driver_sql("CREATE INDEX IF NOT EXISTS ix_product_source_invoice_id ON product (source_invoice_id)")

    for invoice in Invoice.query.filter_by(status="importada").all():
        legacy_items = InvoiceItem.query.filter_by(invoice_id=invoice.id).filter(InvoiceItem.product_id.isnot(None)).all()
        expected = Counter((item.product_id, item.quantity) for item in legacy_items if item.product)
        if not expected:
            continue
        window_end = invoice.imported_at + timedelta(seconds=30)
        candidates = Movement.query.filter(
            Movement.invoice_id.is_(None),
            Movement.kind == "entrada",
            Movement.responsible_id == invoice.imported_by_id,
            Movement.reference == f"NF-e {invoice.number}",
            Movement.created_at >= invoice.imported_at,
            Movement.created_at <= window_end,
        ).all()
        if Counter((movement.product_id, movement.quantity) for movement in candidates) == expected:
            for movement in candidates:
                movement.invoice_id = invoice.id
    db.session.commit()


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=os.environ.get("FLASK_ENV") != "production")